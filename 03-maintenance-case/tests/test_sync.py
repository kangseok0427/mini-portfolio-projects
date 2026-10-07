from before.legacy import sync as legacy, LegacyFailure
from after.sync import sync, normalize
from after.export import export
from dataset import PAGES
import httpx, pytest, sqlite3
from openpyxl import load_workbook


def failing_api():
    calls = {}

    def fetch(page):
        calls[page] = calls.get(page, 0) + 1
        if page == 2 and calls[page] == 1:
            req = httpx.Request("GET", "https://api.example/orders")
            raise httpx.HTTPStatusError(
                "503", request=req, response=httpx.Response(503, request=req)
            )
        return PAGES.get(page, [])

    return fetch


def test_before_stops_with_partial_duplicate_csv(tmp_path):
    with pytest.raises(LegacyFailure) as e:
        legacy(failing_api(), tmp_path, lambda *_: None)
    assert len(e.value.rows) == 3 and len({r["id"] for r in e.value.rows}) == 2
    assert (tmp_path / "orders.csv").exists()


def test_after_retries_validates_deduplicates_and_exports(tmp_path):
    rows, duplicates, errors, retries = sync(
        failing_api(), tmp_path, lambda *_: None, delay=0
    )
    assert len(rows) == 5 and duplicates == 2 and len(errors) == 1 and retries == 1
    export(rows, tmp_path)
    assert load_workbook(tmp_path / "orders.xlsx").active.max_row == 6
    with sqlite3.connect(tmp_path / "orders.sqlite3") as db:
        assert db.execute("SELECT count(*) FROM orders").fetchone()[0] == 5


def test_terminal_page_failure_does_not_stop_other_pages(tmp_path):
    def fetch(page):
        if page == 2:
            raise TimeoutError()
        return PAGES.get(page, [])

    rows, _, errors, _ = sync(fetch, tmp_path, lambda *_: None, retries=1, delay=0)
    assert len(rows) == 4 and len(errors) == 2


def test_money_is_decimal_and_bad_values_rejected():
    assert normalize({"id": "a", "name": "b", "amount": "0.29"})["amount"] == "0.29"
    with pytest.raises(ValueError):
        normalize({"id": "a", "name": "b", "amount": "NaN"})


def test_malformed_json_page_is_isolated(tmp_path):
    def fetch(page):
        if page == 2:
            raise ValueError("Invalid JSON")
        return PAGES.get(page, [])

    rows, _, errors, retries = sync(fetch, tmp_path, lambda *_: None, delay=0)
    assert len(rows) == 4
    assert any("JSON" in error for error in errors)
    assert retries == 0


def test_export_formula_text_is_safe(tmp_path):
    export([{"id": "a", "name": "=1+1", "amount": "10.00"}], tmp_path)
    assert load_workbook(tmp_path / "orders.xlsx").active["B2"].data_type != "f"
