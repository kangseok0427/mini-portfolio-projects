import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def safe(value):
    text = str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def export(rows, folder):
    with open(folder / "orders.csv", "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["주문번호", "상품명", "금액"])
        writer.writerows([safe(r["id"]), safe(r["name"]), r["amount"]] for r in rows)
    book = Workbook()
    sheet = book.active
    sheet.title = "주문내역"
    sheet.append(["주문번호", "상품명", "금액"])
    for row in rows:
        sheet.append([safe(row["id"]), safe(row["name"]), float(row["amount"])])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173B54")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 38
    sheet.column_dimensions["C"].width = 18
    for row in sheet.iter_rows(min_row=2):
        row[2].number_format = "#,##0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(folder / "orders.xlsx")
