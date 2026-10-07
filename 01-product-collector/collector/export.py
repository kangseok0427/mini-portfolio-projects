from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo


def safe_text(value):
    text = str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def export_excel(products, path, summary):
    book = Workbook()
    sheet = book.active
    sheet.title = "상품정보"
    sheet.append(["상품명", "가격", "통화", "URL", "재고", "평점", "출처"])
    for product in products:
        sheet.append(
            [
                safe_text(product.name),
                product.price,
                safe_text(product.currency),
                safe_text(product.url),
                safe_text(product.availability),
                safe_text(product.rating),
                safe_text(product.source),
            ]
        )
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="173B54")
    for row in sheet.iter_rows(min_row=2):
        row[1].number_format = "#,##0.00"
        if str(row[3].value).startswith(("https://", "http://")):
            row[3].hyperlink = row[3].value
            row[3].style = "Hyperlink"
        for cell in row:
            cell.alignment = Alignment(vertical="top")
    for column, width in {
        "A": 42,
        "B": 15,
        "C": 10,
        "D": 65,
        "E": 20,
        "F": 14,
        "G": 22,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    if products:
        table = Table(displayName="Products", ref=f"A1:G{len(products) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        sheet.add_table(table)
    info = book.create_sheet("실행요약")
    info.append(["항목", "값"])
    for key, value in summary.items():
        info.append([safe_text(key), safe_text(value)])
    info.column_dimensions["A"].width = 25
    info.column_dimensions["B"].width = 90
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
