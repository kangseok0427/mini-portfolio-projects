from openpyxl import load_workbook
from collector.pipeline import collect
from collector.models import Product
from collector.export import export_excel
from collector.network import PublicFetcher
import pytest


def test_demo_deduplicates_and_keeps_valid_rows(tmp_path):
    rows, duplicates, errors = collect(
        "", "demo", 3, tmp_path / "out.xlsx", lambda *_: None
    )
    assert len(rows) == 5 and duplicates == 1 and len(errors) == 1
    book = load_workbook(tmp_path / "out.xlsx")
    assert book["상품정보"].max_row == 6 and book["상품정보"]["B2"].value == 59000
    assert book["상품정보"].freeze_panes == "A2"


def test_keyword_filters(tmp_path):
    rows, _, _ = collect("무선", "demo", 1, tmp_path / "out.xlsx", lambda *_: None)
    assert len(rows) == 2


def test_formula_injection_is_text(tmp_path):
    export_excel(
        [Product('=HYPERLINK("bad")', 5, "KRW", "https://example.com/a")],
        tmp_path / "out.xlsx",
        {},
    )
    cell = load_workbook(tmp_path / "out.xlsx").active["A2"]
    assert cell.data_type != "f" and cell.value.startswith("'")


def test_later_page_survives_failed_page(tmp_path):
    class Fetcher:
        def fetch(self, url):
            if "page-2" not in url:
                raise ValueError("HTTP 503")
            return (
                '<article class="product_pod"><h3><a href="a.html" title="Test"></a></h3><p class="price_color">£9.50</p><p class="availability">In stock</p><p class="star-rating One"></p></article>',
                url,
            )

    rows, _, errors = collect(
        "", "web", 2, tmp_path / "out.xlsx", lambda *_: None, Fetcher()
    )
    assert len(rows) == 1 and len(errors) == 1


def test_unapproved_host_is_rejected():
    with pytest.raises(ValueError):
        PublicFetcher().validate("https://example.com/")


def test_private_ip_is_rejected(monkeypatch):
    monkeypatch.setenv("ALLOWED_HOSTS", "localhost")
    with pytest.raises(ValueError):
        PublicFetcher().validate("http://localhost/")
