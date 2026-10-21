"""Exercise real local apps and record their browser UI. No responses are mocked.

Run from the repository root: python tools/record_demos.py
Requires Chrome and: python -m playwright install ffmpeg
Only the apps' documented sample modes are used; no paid LLM calls are made.
"""
import asyncio
from contextlib import contextmanager
import json
import os
from pathlib import Path
import subprocess
import sys
import time
import httpx
import imageio_ffmpeg
from openpyxl import load_workbook
from playwright.async_api import async_playwright

ROOT = Path(__file__).resolve().parents[1]
PROJECTS = ['01-product-collector', '02-document-assistant', '03-maintenance-case']


@contextmanager
def servers():
    processes, logs = [], []
    try:
        for port, name in enumerate(PROJECTS, 8101):
            folder = ROOT / name
            logfile = open(ROOT / 'tmp' / f'{port}.log', 'w', encoding='utf-8')
            logs.append(logfile)
            env = dict(os.environ, LLM_PROVIDER='demo', PYTHONUTF8='1')
            processes.append(subprocess.Popen([sys.executable, '-m', 'uvicorn', 'app:app',
                '--host', '127.0.0.1', '--port', str(port)], cwd=folder, env=env,
                stdout=logfile, stderr=logfile, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0)))
        with httpx.Client(trust_env=False) as client:
            for port in range(8101, 8104):
                for attempt in range(120):
                    try:
                        response = client.get(f'http://127.0.0.1:{port}/', timeout=1)
                        response.raise_for_status()
                        break
                    except httpx.HTTPError:
                        time.sleep(.5)
                else:
                    raise RuntimeError(f'App on port {port} did not start; inspect tmp/{port}.log')
        yield
    finally:
        for process in processes:
            process.terminate()
        for process in processes:
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()
        for logfile in logs:
            logfile.close()


async def caption(page, text):
    # Caption labels the mode; it never replaces app results or network responses.
    await page.evaluate('''text => {
        let el = document.getElementById('recording-caption');
        if(!el) { el=document.createElement('div');el.id='recording-caption';document.body.append(el); }
        el.textContent=text;
        el.style.cssText='position:fixed;bottom:16px;left:24px;right:24px;padding:14px 20px;background:#142334;color:white;border:1px solid #59728b;border-radius:12px;z-index:99999;font:600 18px sans-serif;box-shadow:0 4px 24px #0004';
    }''', text)


async def record(browser, index):
    name, port = PROJECTS[index], 8101 + index
    folder = ROOT / name / 'demo'
    context = await browser.new_context(viewport={'width':1280, 'height':900},
        device_scale_factor=1, record_video_dir=str(ROOT / 'tmp' / 'video'),
        record_video_size={'width':1280, 'height':900}, locale='ko-KR')
    page = await context.new_page()
    errors = []
    page.on('pageerror', lambda err: errors.append(str(err)))
    evidence = {'mode':'local sample; actual browser/API execution', 'project':name}
    await page.goto(f'http://127.0.0.1:{port}/')
    await page.wait_for_timeout(1500)
    if index == 0:
        await caption(page, '실제 실행 · 자체 샘플 상품 → 중복 제거 → Excel 파일')
        await page.screenshot(path=str(folder/'01-start.png'))
        await page.wait_for_timeout(3500)
        await page.click('#start')
        await page.wait_for_selector('#results:not(.hidden)')
        assert await page.locator('#count').inner_text() == '5'
        assert await page.locator('#duplicates').inner_text() == '1'
        await page.locator('#results').scroll_into_view_if_needed()
        await caption(page, '정상 상품 5개 · 중복 1개 제거 · 잘못된 상품 1개 건너뜀')
        await page.wait_for_timeout(5500)
        await page.screenshot(path=str(folder/'02-results.png'))
        async with page.expect_download() as download_info:
            await page.click('#download')
        download = await download_info.value
        saved = ROOT/'tmp'/'products.xlsx'
        await download.save_as(saved)
        workbook = load_workbook(saved)
        assert workbook.active.max_row == 6
        assert workbook.active.freeze_panes == 'A2'
        evidence.update(rows=5, duplicates=1, excel_rows=workbook.active.max_row)
        await caption(page, 'Excel 다운로드 성공 · 실제 파일의 5개 상품 행과 헤더 고정 검증')
        await page.wait_for_timeout(6000)
        await page.screenshot(path=str(folder/'03-export.png'))
    elif index == 1:
        await caption(page, '실제 실행 · 문서 업로드와 출처 검색 · LLM 호출 없는 검색 모드')
        await page.screenshot(path=str(folder/'01-start.png'))
        await page.wait_for_timeout(3500)
        docs = sorted((ROOT/name/'demo/documents').glob('*'))
        await page.set_input_files('#files', [str(p) for p in docs])
        await page.click('#upload')
        await page.wait_for_function("document.querySelectorAll('#documents .citation').length === 3")
        await caption(page, '샘플 TXT 2개 + PDF 1개 업로드 · 문서 3개 색인')
        await page.wait_for_timeout(4500)
        await page.screenshot(path=str(folder/'02-upload.png'))
        await page.fill('#question', '계약 기간과 총 계약금은 얼마인가요?')
        await page.wait_for_timeout(2500)
        await page.click('#ask')
        await page.wait_for_selector('#chat .assistant .citation')
        await page.locator('#chat .assistant .citation').first.locator('summary').click()
        await page.locator('#chat .assistant').scroll_into_view_if_needed()
        await caption(page, '질문 관련 원문과 파일명·페이지 표시 · 생성형 AI 답변과 구분')
        await page.wait_for_timeout(6500)
        await page.screenshot(path=str(folder/'03-sources.png'))
        evidence.update(documents=3, citations=await page.locator('#chat .assistant .citation').count())
    else:
        await caption(page, '실제 실행 · 자체 제작 장애 재현 · 동일 주문 API의 BEFORE / AFTER')
        await page.screenshot(path=str(folder/'01-start.png'))
        await page.wait_for_timeout(3500)
        await page.click('#before')
        await page.wait_for_function("document.querySelector('#before-count').textContent.trim() === '3'")
        await caption(page, 'BEFORE · API 503에서 작업 중단 · 중복을 포함한 3개 행')
        await page.wait_for_timeout(5000)
        await page.screenshot(path=str(folder/'02-before.png'))
        await page.click('#after')
        await page.wait_for_function("document.querySelector('#after-count').textContent.trim() === '5'")
        assert await page.locator('#after-duplicates').inner_text() == '2'
        assert await page.locator('#after-retries').inner_text() == '1'
        await caption(page, 'AFTER · 정상 주문 5개 · 중복 2개 제거 · API 재시도 1회')
        await page.wait_for_timeout(6500)
        await page.screenshot(path=str(folder/'03-after.png'))
        async with page.expect_download() as download_info:
            await page.locator('#after-files a').filter(has_text='.xlsx').click()
        download = await download_info.value
        saved = ROOT/'tmp'/'orders.xlsx'
        await download.save_as(saved)
        assert load_workbook(saved).active.max_row == 6
        await caption(page, 'CSV · Excel · SQLite 저장 · 실제 Excel 파일 5개 주문 검증')
        await page.wait_for_timeout(5000)
        evidence.update(before_rows=3, after_rows=5, removed_duplicates=2, retries=1, excel_rows=6)
    assert not errors, errors
    video = page.video
    await context.close()
    raw = Path(await video.path())
    target = folder/'demo.mp4'
    subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), '-y','-i',str(raw),'-an',
        '-c:v','libx264','-crf','25','-preset','medium','-pix_fmt','yuv420p',
        '-movflags','+faststart',str(target)], check=True, capture_output=True)
    assert target.stat().st_size < 10_000_000
    evidence.update(video_bytes=target.stat().st_size, browser_errors=errors)
    (folder/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(evidence,ensure_ascii=True), flush=True)


async def main():
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(channel='chrome', headless=True)
        try:
            for index in range(3):
                await record(browser,index)
        finally:
            await browser.close()


if __name__ == '__main__':
    (ROOT/'tmp').mkdir(exist_ok=True)
    with servers():
        asyncio.run(main())
